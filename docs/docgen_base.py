import os
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
VERDE = RGBColor(0x2E, 0x7D, 0x32)

d = Document()
_fig = [0]
_tab = [0]


def _fuente(estilo, nombre, tam, color=None, negrita=None):
    f = estilo.font
    f.name = nombre
    f.size = Pt(tam)
    if color is not None:
        f.color.rgb = color
    if negrita is not None:
        f.bold = negrita
    rpr = estilo.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.append(rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(a), nombre)


def configurar():
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin = sec.bottom_margin = Cm(2.3)
    st = d.styles
    _fuente(st['Normal'], 'Calibri', 11)
    st['Normal'].paragraph_format.space_after = Pt(6)
    st['Normal'].paragraph_format.line_spacing = 1.15
    _fuente(st['Heading 1'], 'Calibri', 18, VERDE, True)
    _fuente(st['Heading 2'], 'Calibri', 14, RGBColor(0x1B, 0x5E, 0x20), True)
    _fuente(st['Heading 3'], 'Calibri', 12, RGBColor(0x33, 0x33, 0x33), True)
    st['Heading 1'].paragraph_format.space_before = Pt(0)
    st['Heading 1'].paragraph_format.space_after = Pt(10)
    st['Heading 2'].paragraph_format.space_before = Pt(14)
    st['Heading 3'].paragraph_format.space_before = Pt(10)
    for nombre in ('Caption',):
        _fuente(st[nombre], 'Calibri', 9.5, RGBColor(0x45, 0x5A, 0x64), False)
        st[nombre].font.italic = True


def _campo(par, instr):
    r = par.add_run()
    for tipo, txt in (('begin', None), (None, instr), ('separate', None), (None, '1'), ('end', None)):
        if tipo:
            e = OxmlElement('w:fldChar')
            e.set(qn('w:fldCharType'), tipo)
        elif txt == instr:
            e = OxmlElement('w:instrText')
            e.set(qn('xml:space'), 'preserve')
            e.text = instr
        else:
            e = OxmlElement('w:t')
            e.text = txt
        r._r.append(e)


def pie_y_encabezado(titulo_corto):
    sec = d.sections[0]
    sec.different_first_page_header_footer = True
    hp = sec.header.paragraphs[0]
    hp.text = titulo_corto
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in hp.runs:
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x75, 0x75, 0x75)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run('Página ').font.size = Pt(9)
    _campo(fp, 'PAGE')
    for r in fp.runs:
        r.font.size = Pt(9)


def H(texto, nivel=1, salto=True):
    if nivel == 1 and salto:
        d.add_page_break()
    return d.add_heading(texto, nivel)


def P(texto, negrita_inicio=None, centrado=False, cursiva=False):
    par = d.add_paragraph()
    if negrita_inicio:
        par.add_run(negrita_inicio).bold = True
    r = par.add_run(texto)
    r.italic = cursiva
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER if centrado else WD_ALIGN_PARAGRAPH.JUSTIFY
    return par


def B(items, estilo='List Bullet'):
    for i in items:
        par = d.add_paragraph(style=estilo)
        if isinstance(i, tuple):
            par.add_run(i[0]).bold = True
            par.add_run(i[1])
        else:
            par.add_run(i)
        par.paragraph_format.space_after = Pt(2)


def N(items):
    B(items, 'List Number')


def _sombrear(celda, hexa):
    tcpr = celda._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:color'), 'auto')
    sh.set(qn('w:fill'), hexa)
    tcpr.append(sh)


def T(encabezado, filas, anchos=None, titulo=None, tam=9):
    if titulo:
        _tab[0] += 1
        c = d.add_paragraph(f'Tabla {_tab[0]}. {titulo}', style='Caption')
        c.paragraph_format.keep_with_next = True
    t = d.add_table(rows=1, cols=len(encabezado))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, x in enumerate(encabezado):
        cel = t.rows[0].cells[i]
        cel.text = ''
        run = cel.paragraphs[0].add_run(x)
        run.bold = True
        run.font.size = Pt(tam)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _sombrear(cel, '2E7D32')
    trpr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement('w:tblHeader')
    th.set(qn('w:val'), 'true')
    trpr.append(th)
    for n, fila in enumerate(filas):
        cs = t.add_row().cells
        for i, x in enumerate(fila):
            cs[i].text = ''
            run = cs[i].paragraphs[0].add_run(str(x))
            run.font.size = Pt(tam)
            cs[i].paragraphs[0].paragraph_format.space_after = Pt(0)
            if n % 2:
                _sombrear(cs[i], 'F1F8E9')
    if anchos:
        for fila in t.rows:
            for i, w in enumerate(anchos):
                fila.cells[i].width = Cm(w)
    d.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def CODE(texto, tam=7.5):
    for linea in texto.split('\n'):
        par = d.add_paragraph()
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.line_spacing = 1.0
        par.paragraph_format.left_indent = Cm(0.3)
        r = par.add_run(linea if linea else ' ')
        r.font.name = 'Consolas'
        r._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Consolas')
        r.font.size = Pt(tam)
        pPr = par._p.get_or_add_pPr()
        sh = OxmlElement('w:shd')
        sh.set(qn('w:val'), 'clear')
        sh.set(qn('w:color'), 'auto')
        sh.set(qn('w:fill'), 'F5F5F5')
        pPr.append(sh)
    d.add_paragraph().paragraph_format.space_after = Pt(2)


def FIG(ruta, pie, ancho=15.5):
    ruta = os.path.join(RAIZ, ruta)
    par = d.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.keep_with_next = True
    par.add_run().add_picture(ruta, width=Cm(ancho))
    _fig[0] += 1
    c = d.add_paragraph(f'Figura {_fig[0]}. {pie}', style='Caption')
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER


def FIG_ALTO(ruta, pie, alto=19.0):
    ruta = os.path.join(RAIZ, ruta)
    par = d.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.keep_with_next = True
    par.add_run().add_picture(ruta, height=Cm(alto))
    _fig[0] += 1
    c = d.add_paragraph(f'Figura {_fig[0]}. {pie}', style='Caption')
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER


def LEER(rel):
    with open(os.path.join(RAIZ, rel), encoding='utf-8') as f:
        return f.read().rstrip('\n').replace('\t', '    ')


def TOC():
    par = d.add_paragraph()
    r = par.add_run()
    for tipo in ('begin', 'instr', 'separate', 'texto', 'end'):
        if tipo == 'instr':
            e = OxmlElement('w:instrText')
            e.set(qn('xml:space'), 'preserve')
            e.text = 'TOC \\o "1-2" \\h \\z \\u'
        elif tipo == 'texto':
            e = OxmlElement('w:t')
            e.text = 'Actualice el campo para ver el índice.'
        else:
            e = OxmlElement('w:fldChar')
            e.set(qn('w:fldCharType'), tipo)
        r._r.append(e)


def portada(titulo, subtitulo, lineas):
    for _ in range(5):
        d.add_paragraph()
    t = d.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(titulo)
    r.bold = True
    r.font.size = Pt(34)
    r.font.color.rgb = VERDE
    s = d.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s.add_run(subtitulo)
    r.font.size = Pt(16)
    for _ in range(6):
        d.add_paragraph()
    for l in lineas:
        q = d.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        q.add_run(l).font.size = Pt(13)
